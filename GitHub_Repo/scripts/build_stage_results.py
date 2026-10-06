"""
build_stage_results.py -- builds the LLM-judge results workbook and figures
for the three policy stages (A, B, C).

Produces the same table and figure set as build_results.py, but reads the
judge output CSVs written by run_llm_judge.py instead of the manually coded
sheet. Each stage gets its own Tables and Figures sheets, plus a cross-stage
summary sheet. Standalone PNGs are saved to figs_stages/.

Usage (from the repository root):
    python3 scripts/build_stage_results.py \
        policy/judge_output_full_coding_policy_A.csv \
        policy/judge_output_full_coding_policy_B.csv \
        policy/judge_output_full_coding_policy_C.csv \
        stage_results.xlsx --labels "Names only" "Full definitions" "Definitions+rules"

Without --labels the stages are labelled A, B, C in the order given.
"""

import sys
import os
import re
import csv as csvmod
from collections import Counter, defaultdict

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np

try:
    from wordcloud import WordCloud, STOPWORDS
    HAVE_WORDCLOUD = True
except ImportError:
    HAVE_WORDCLOUD = False

# ---------------------------------------------------------------------------
# STYLE (matches build_results.py)
# ---------------------------------------------------------------------------

AXES_BG = "#EAEAF2"
GRID_COLOR = "#FFFFFF"
PALETTE = ["#3B0F70", "#8C2981", "#B83280", "#DE4968", "#F2A65A"]  # purple-pink-orange, magma-family
SINGLE_BAR_COLOR = "#2E5EAA"  # navy, for single-series bars
HEATMAP_CMAP = "magma"
WORDCLOUD_COLORS = ["#3B0F70", "#8C2981", "#B83280", "#DE4968", "#F2A65A"]

plt.rcParams.update({
    "axes.facecolor": AXES_BG, "axes.edgecolor": AXES_BG, "axes.grid": True,
    "grid.color": GRID_COLOR, "grid.linewidth": 1.3, "axes.axisbelow": True,
    "figure.facecolor": "white", "font.family": "sans-serif", "text.color": "#2B2B2B",
    "axes.labelcolor": "#2B2B2B", "xtick.color": "#2B2B2B", "ytick.color": "#2B2B2B",
})

HEADER_FONT = Font(bold=True, size=11, color="000000")
ROW_LABEL_FONT = Font(bold=True, size=11, color="000000")
TITLE_FONT = Font(bold=True, size=13, color="000000")
NOTE_FONT = Font(italic=True, size=10, color="444444")
THIN = Side(style="thin", color="AAAAAA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
GREEN_SCALE = ColorScaleRule(
    start_type="min", start_color="FFFFFF",
    mid_type="percentile", mid_value=50, mid_color="C6E8B0",
    end_type="max", end_color="4E8A3C",
)

CATEGORY_ORDER = [
    "Blaming", "Economic abuse", "Emotional abuse", "Intimidation",
    "Isolation", "Minimising/denying", "Physical abuse", "Sexual abuse", "Threats",
]
PERSONA_ORDER = ["baseline", "girlfriend", "boyfriend", "combined"]
REAL_PERSONAS = ["baseline", "girlfriend", "boyfriend"]
VALID_L2 = {"Harmful", "Non-harmful"}
VALID_ROLE = {"Facilitator", "Enabler"}
NA_VARIANTS = {"n/a", "na", "not applicable", "non applicable", "none", ""}
FILLER_WORDS = {
    "hey", "please", "trying", "listen", "just", "really", "im", "youre",
    "dont", "know", "want", "think", "feel", "feeling", "like", "okay",
    "ok", "going", "get", "got", "would", "could", "way", "thing",
    "things", "sorry", "understand", "ive", "youve", "id",
    "theyre", "arent", "isnt", "didnt", "cant", "wont", "wasnt",
    "youll", "ill", "were", "theyve", "weve", "whats", "thats", "its",
    "us", "let", "lets", "look", "looking", "looks", "take", "taking",
    "come", "coming", "much", "make", "making", "makes", "something",
    "actually", "instead", "even", "right", "now", "need", "needs",
    "say", "saying", "says", "see", "seeing", "little", "always", "back",
    "one", "certainly", "perhaps", "sound", "sounds", "moment", "second",
    "exactly", "sure", "still", "yet", "also", "well", "maybe",
}
FIG_DIR = "figs_stages"
if os.path.isdir(FIG_DIR):
    import shutil
    shutil.rmtree(FIG_DIR)
os.makedirs(FIG_DIR, exist_ok=True)

COL_PERSONA, COL_SCENARIO, COL_CAT = "persona_condition", "scenario_id", "Abuse_type"
COL_L1, COL_L2, COL_SUBCAT, COL_ROLE = ("llm_layer1_response_type", "llm_layer2_safety_judgment",
                                          "llm_layer3_transgression_subcat", "llm_layer3_ai_role")
COL_PARSE_ERROR = "llm_parse_error"


# ---------------------------------------------------------------------------
# STYLING HELPERS
# ---------------------------------------------------------------------------

def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = CENTER


def style_data_cell(cell, center=True):
    cell.border = BORDER
    cell.alignment = CENTER if center else Alignment(horizontal="left", vertical="center")


def autosize(ws, min_width=10, max_width=55):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            col = cell.column_letter
            widths[col] = max(widths.get(col, min_width), min(max_width, len(str(cell.value)) + 2))
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def apply_green_scale(ws, first_col, last_col, first_row, last_row):
    rng = f"{get_column_letter(first_col)}{first_row}:{get_column_letter(last_col)}{last_row}"
    ws.conditional_formatting.add(rng, GREEN_SCALE)


def write_table(ws, headers, rows, start_row=1, title=None):
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = TITLE_FONT
        r += 1
    header_row = r
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    style_header_row(ws, header_row, len(headers))
    r += 1
    for row_vals in rows:
        for c, v in enumerate(row_vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            style_data_cell(cell, center=(c > 1))
            if c == 1:
                cell.font = ROW_LABEL_FONT
        r += 1
    return r + 1


def write_table_heatmap(ws, headers, rows, start_row=1, title=None, label_cols=1):
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = TITLE_FONT
        r += 1
    header_row = r
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    style_header_row(ws, header_row, len(headers))
    r += 1
    data_start = r
    for row_vals in rows:
        for c, v in enumerate(row_vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            style_data_cell(cell, center=(c > 1))
            if c == 1:
                cell.font = ROW_LABEL_FONT
        r += 1
    data_end = r - 1
    if data_end >= data_start:
        apply_green_scale(ws, label_cols + 1, len(headers), data_start, data_end)
    return r + 1


# ---------------------------------------------------------------------------
# CHART BUILDERS -- individual PNGs, matching build_results.py
# ---------------------------------------------------------------------------

def save_fig(fig, name):
    path = os.path.join(FIG_DIR, f"{name}.png")
    fig.savefig(path, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def bar_chart_single(categories, values, title, ylabel, name, ylim=None):
    fig, ax = plt.subplots(figsize=(7.5, 4.5))
    ax.bar(categories, values, color=SINGLE_BAR_COLOR, width=0.6)
    for i, v in enumerate(values):
        ax.text(i, v + (max(values) * 0.02 if max(values) else 0.5), f"{v:.1f}",
                 ha="center", va="bottom", fontsize=9)
    ax.set_title(title, fontsize=12, fontweight="bold", pad=10)
    ax.set_ylabel(ylabel, fontsize=10)
    if ylim:
        ax.set_ylim(*ylim)
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(axis="x", rotation=20)
    plt.tight_layout()
    return save_fig(fig, name)


def bar_chart_grouped(categories, series_dict, title, ylabel, name):
    fig, ax = plt.subplots(figsize=(9.5, 5.2))
    n = len(series_dict)
    x = np.arange(len(categories))
    width = 0.8 / n
    for i, (label, vals) in enumerate(series_dict.items()):
        ax.bar(x + i * width - 0.4 + width / 2, vals, width, label=label, color=PALETTE[i % len(PALETTE)])
    ax.set_xticks(x)
    ax.set_xticklabels(categories, rotation=35, ha="right", fontsize=9)
    ax.set_title(title, fontsize=12, fontweight="bold", pad=10)
    ax.set_ylabel(ylabel, fontsize=10)
    ax.legend(frameon=False, fontsize=9, loc="upper center", bbox_to_anchor=(0.5, -0.22), ncol=n)
    ax.spines[["top", "right"]].set_visible(False)
    plt.tight_layout()
    return save_fig(fig, name)


def stacked_bar(categories, series_dict, title, ylabel, name):
    fig, ax = plt.subplots(figsize=(7, 5))
    bottom = np.zeros(len(categories))
    for i, (label, vals) in enumerate(series_dict.items()):
        vals = np.array(vals, dtype=float)
        ax.bar(categories, vals, bottom=bottom, label=label, color=PALETTE[i % len(PALETTE)], width=0.6)
        bottom += vals
    ax.set_title(title, fontsize=12, fontweight="bold", pad=10)
    ax.set_ylabel(ylabel, fontsize=10)
    ax.legend(frameon=False, fontsize=9, loc="upper center", bbox_to_anchor=(0.5, -0.15), ncol=len(series_dict))
    ax.spines[["top", "right"]].set_visible(False)
    plt.tight_layout()
    return save_fig(fig, name)


def heatmap_chart(categories, response_types, matrix, title, name):
    fig, ax = plt.subplots(figsize=(6.5, 5.8))
    ax.grid(False)
    ax.set_facecolor("white")
    im = ax.imshow(matrix, cmap=HEATMAP_CMAP, vmin=0, vmax=100, aspect="auto")
    ax.set_xticks(range(len(response_types))); ax.set_xticklabels(response_types, fontsize=9)
    ax.set_yticks(range(len(categories))); ax.set_yticklabels(categories, fontsize=9)
    for i in range(len(categories)):
        for j in range(len(response_types)):
            val = matrix[i][j]
            ax.text(j, i, f"{val:.0f}", ha="center", va="center", fontsize=8,
                     color="white" if val < 50 else "black")
    ax.set_title(title, fontsize=12, fontweight="bold", pad=10)
    cbar = fig.colorbar(im, ax=ax, fraction=0.046, pad=0.04)
    cbar.set_label("% of category's trials", fontsize=9)
    plt.tight_layout()
    return save_fig(fig, name)


def _muted_color_func(word=None, font_size=None, position=None, orientation=None,
                        font_path=None, random_state=None):
    return WORDCLOUD_COLORS[(random_state.randint(0, 10**6) if random_state else 0) % len(WORDCLOUD_COLORS)]


def wordcloud_image(text, title, name):
    if not HAVE_WORDCLOUD or not text.strip():
        return None
    stop = set(STOPWORDS) | FILLER_WORDS
    wc = WordCloud(width=900, height=520, background_color="white", stopwords=stop,
                    max_words=50, prefer_horizontal=0.95, color_func=_muted_color_func).generate(text)
    fig, ax = plt.subplots(figsize=(7.5, 4.3))
    ax.imshow(wc, interpolation="bilinear")
    ax.axis("off")
    ax.set_title(title, fontsize=12, fontweight="bold", pad=10)
    plt.tight_layout()
    return save_fig(fig, name)


def plain_table_image(row_labels, col_labels, matrix, title, name, value_suffix="%", figsize=None):
    """Plain white/bordered table with no colour shading, for use as a
    table in the Results chapter."""
    n_rows, n_cols = len(row_labels), len(col_labels)
    max_label_len = max(len(str(rl)) for rl in row_labels)
    label_units = max(1.1, max_label_len * 0.105)
    if figsize is None:
        figsize = (1.3 * n_cols + label_units + 0.8, 0.5 * n_rows + 1.3)
    fig, ax = plt.subplots(figsize=figsize)
    ax.set_xlim(0, n_cols + label_units)
    ax.set_ylim(0, n_rows + 1)
    ax.invert_yaxis()
    ax.axis("off")
    ax.set_facecolor("white")

    for j, col in enumerate(col_labels):
        ax.text(label_units + j + 0.5, 0.5, col, ha="center", va="center", fontsize=10, fontweight="bold")
    ax.plot([label_units, n_cols + label_units], [1, 1], color="#333333", linewidth=1.3)

    for i, row_label in enumerate(row_labels):
        ax.text(label_units - 0.15, i + 1.5, row_label, ha="right", va="center", fontsize=10, fontweight="bold")
        for j, val in enumerate(matrix[i]):
            ax.add_patch(plt.Rectangle((label_units + j, i + 1), 1, 1, facecolor="white",
                                         edgecolor="#CCCCCC", linewidth=0.8))
            label = f"{val:.1f}{value_suffix}"
            ax.text(label_units + j + 0.5, i + 1.5, label, ha="center", va="center", fontsize=9.5, color="black")

    ax.set_title(title, fontsize=13, fontweight="bold", pad=14, loc="left")
    plt.tight_layout()
    return save_fig(fig, name)


def clean_text(raw):
    if not raw:
        return ""
    text = raw.replace("'", "").replace("\u2019", "")
    return re.sub(r"[^a-zA-Z\s]", " ", text)


def normalize_subcat(s):
    if s is None:
        return "N/A"
    s_clean = str(s).strip()
    return "N/A" if s_clean.lower() in NA_VARIANTS else s_clean


def primary_code(l1):
    if not l1:
        return None
    return re.split(r"\s+and\s+", str(l1).strip())[0].strip()


# ---------------------------------------------------------------------------
# LOAD ONE STAGE'S DATA
# ---------------------------------------------------------------------------

def load_stage(path):
    with open(path, "r", encoding="utf-8") as f:
        rows = list(csvmod.DictReader(f))
    n_before = len(rows)
    rows = [r for r in rows if not (r.get(COL_PARSE_ERROR) or "").strip()]
    dropped = n_before - len(rows)
    data = []
    for r in rows:
        data.append({
            "persona": r.get(COL_PERSONA), "scenario": r.get(COL_SCENARIO), "category": r.get(COL_CAT),
            "raw_text": r.get("raw_response_text", ""),
            "l1": r.get(COL_L1), "l2": r.get(COL_L2),
            "role": r.get(COL_ROLE), "subcat": normalize_subcat(r.get(COL_SUBCAT)),
        })
    out = list(data)
    for row in data:
        if row["persona"] in ("girlfriend", "boyfriend"):
            new_row = dict(row); new_row["persona"] = "combined"
            out.append(new_row)
    return out, dropped


# ---------------------------------------------------------------------------
# ANALYSIS (same logic as build_results.py, condensed)
# ---------------------------------------------------------------------------

def layer1_by_persona(data):
    rows = []
    for p in PERSONA_ORDER:
        sub = [r for r in data if r["persona"] == p]
        n = len(sub)
        codes = Counter(primary_code(r["l1"]) for r in sub)
        rows.append([p, n] + [round(100 * codes.get(k, 0) / n, 1) if n else 0
                                for k in ["SRM", "NFP", "RD", "R/BK"]])
    return rows


def layer2_by_persona(data):
    rows = []
    for p in PERSONA_ORDER:
        sub = [r for r in data if r["persona"] == p]
        valid = [r for r in sub if r["l2"] in VALID_L2]
        n = len(valid)
        h = sum(1 for r in valid if r["l2"] == "Harmful")
        rows.append([p, n, len(sub) - n, round(100 * h / n, 1) if n else 0,
                      round(100 * (n - h) / n, 1) if n else 0])
    return rows


def rate_by_category(data, target_kind):
    rows = []
    for cat in CATEGORY_ORDER:
        row = [cat]
        for p in PERSONA_ORDER:
            sub = [r for r in data if r["persona"] == p and r["category"] == cat]
            if target_kind == "harmful":
                valid = [r for r in sub if r["l2"] in VALID_L2]
                n = len(valid)
                pct = 100 * sum(1 for r in valid if r["l2"] == "Harmful") / n if n else 0
            else:
                n = len(sub)
                pct = 100 * sum(1 for r in sub if primary_code(r["l1"]) == target_kind) / n if n else 0
            row.append(round(pct, 1))
        rows.append(row)
    return rows


def ai_role_by_persona(data):
    rows = []
    for p in PERSONA_ORDER:
        sub = [r for r in data if r["persona"] == p and r["l2"] == "Harmful"]
        n = len(sub)
        c = Counter(r["role"] for r in sub if r["role"] in VALID_ROLE)
        rows.append([p, n, round(100 * c.get("Facilitator", 0) / n, 1) if n else 0,
                      round(100 * c.get("Enabler", 0) / n, 1) if n else 0])
    return rows


def subcategory_by_persona(data):
    rows = []
    for p in PERSONA_ORDER:
        sub = [r for r in data if r["persona"] == p and r["l2"] == "Harmful"]
        n_harm = len(sub)
        subcoded = [r for r in sub if r["subcat"] != "N/A"]
        n_sub = len(subcoded)
        counts = Counter(r["subcat"] for r in subcoded)
        if not counts:
            rows.append([p, n_harm, n_sub, "-", 0, 0])
            continue
        for subcat, cnt in sorted(counts.items(), key=lambda kv: -kv[1]):
            rows.append([p, n_harm, n_sub, subcat, cnt, round(100 * cnt / n_sub, 1) if n_sub else 0])
    return rows


# ---------------------------------------------------------------------------
# BUILD ONE STAGE'S SHEETS + FIGURES
# ---------------------------------------------------------------------------

def build_stage(wb, stage_label, data, dropped, fig_start_n):
    prefix = re.sub(r"[^A-Za-z0-9]", "", stage_label)[:20]
    fig_n = fig_start_n

    # ---- Tables (Excel sheet, green conditional-formatting scale) ----
    ws = wb.create_sheet(f"{prefix}_Tables")
    r = 1
    ws.cell(row=r, column=1, value=f"Stage: {stage_label}  ({len(data) - sum(1 for d in data if d['persona']=='combined')} rows, "
                                     f"{dropped} dropped for parse error)").font = TITLE_FONT
    r += 2
    r = write_table_heatmap(ws, ["Persona", "n valid", "n excluded", "Harmful %", "Non-harmful %"],
                              layer2_by_persona(data), start_row=r, title="Layer 2: safety judgment")
    l1_rows = layer1_by_persona(data)
    r = write_table(ws, ["Persona", "n", "SRM %", "NFP %", "RD %", "R/BK %"], l1_rows, start_row=r,
                     title="Layer 1: response type (primary code)")
    harm_rows = rate_by_category(data, "harmful")
    r = write_table_heatmap(ws, ["Abuse category"] + PERSONA_ORDER, harm_rows, start_row=r,
                              title="Harmful rate by category")
    role_rows = ai_role_by_persona(data)
    r = write_table_heatmap(ws, ["Persona", "n Harmful", "Facilitator %", "Enabler %"], role_rows, start_row=r,
                              title="AI role among Harmful responses")
    subcat_rows = subcategory_by_persona(data)
    write_table(ws, ["Persona", "n Harmful", "n subcategory-coded", "Subcategory", "Count", "% of subcoded"],
                subcat_rows, start_row=r, title="Harm subcategory among Harmful responses")
    autosize(ws)

    # ---- Figures (individual PNGs) ----
    ws = wb.create_sheet(f"{prefix}_Figures")
    row_cursor = 1

    def embed(path):
        nonlocal row_cursor
        img = XLImage(path); img.anchor = f"A{row_cursor}"; ws.add_image(img)
        row_cursor += 24

    layer2_rows = layer2_by_persona(data)
    cats = CATEGORY_ORDER

    embed(bar_chart_single(PERSONA_ORDER, [row[3] for row in layer2_rows],
        f"Overall Harmful Rate by Condition ({stage_label})", "Harmful %",
        f"fig{fig_n}_{prefix}_overall_harmful", ylim=(0, 100)))
    fig_n += 1

    embed(bar_chart_grouped(cats, {p: [row[i + 1] for row in harm_rows] for i, p in enumerate(REAL_PERSONAS)},
        f"Harmful Rate by Category ({stage_label})", "Harmful %", f"fig{fig_n}_{prefix}_harmful_by_category"))
    fig_n += 1

    embed(bar_chart_grouped(cats, {p: [row[i + 2] for row in harm_rows] for i, p in enumerate(["girlfriend", "boyfriend"])},
        f"Harmful Rate by Category: Girlfriend vs Boyfriend ({stage_label})", "Harmful %",
        f"fig{fig_n}_{prefix}_harmful_gf_vs_bf"))
    fig_n += 1

    embed(stacked_bar([row[0] for row in l1_rows],
        {"SRM": [row[2] for row in l1_rows], "NFP": [row[3] for row in l1_rows],
         "RD": [row[4] for row in l1_rows], "R/BK": [row[5] for row in l1_rows]},
        f"Response-Type Composition by Persona ({stage_label})", "% of trials",
        f"fig{fig_n}_{prefix}_response_type_composition"))
    fig_n += 1

    embed(stacked_bar([row[0] for row in role_rows],
        {"Facilitator": [row[2] for row in role_rows], "Enabler": [row[3] for row in role_rows]},
        f"AI Role Among Harmful Responses ({stage_label})", "% of Harmful trials",
        f"fig{fig_n}_{prefix}_ai_role"))
    fig_n += 1

    subcat_by_persona_pct = defaultdict(dict)
    for row in subcat_rows:
        persona, n_harm, n_sub, subcat, cnt, pct = row
        if subcat != "-":
            subcat_by_persona_pct[persona][subcat] = pct
    all_subcats = sorted({sc for d in subcat_by_persona_pct.values() for sc in d})
    if all_subcats:
        embed(stacked_bar(PERSONA_ORDER,
            {sc: [subcat_by_persona_pct[p].get(sc, 0) for p in PERSONA_ORDER] for sc in all_subcats},
            f"Harm Subcategory Among Subcategory-Coded Harmful Responses ({stage_label})",
            "% of subcategory-coded responses", f"fig{fig_n}_{prefix}_subcategory"))
        fig_n += 1

    # Plain (uncoloured) SRM/R-BK tables -- not duplicated by any bar chart above
    srm_rows = rate_by_category(data, "SRM")
    rbk_rows = rate_by_category(data, "R/BK")
    embed(plain_table_image(cats, PERSONA_ORDER, [row[1:] for row in srm_rows],
        f"SRM Rate by Category ({stage_label})", f"table{fig_n}_{prefix}_srm_by_category"))
    fig_n += 1
    embed(plain_table_image(cats, PERSONA_ORDER, [row[1:] for row in rbk_rows],
        f"R/BK Rate by Category ({stage_label})", f"table{fig_n}_{prefix}_rbk_by_category"))
    fig_n += 1

    response_types = ["SRM", "NFP", "RD", "R/BK"]
    for p in PERSONA_ORDER:
        matrix = []
        for cat in CATEGORY_ORDER:
            sub = [r for r in data if r["persona"] == p and r["category"] == cat]
            n = len(sub)
            codes = Counter(primary_code(r["l1"]) for r in sub)
            matrix.append([100 * codes.get(rt, 0) / n if n else 0 for rt in response_types])
        embed(heatmap_chart(CATEGORY_ORDER, response_types, matrix,
            f"Response Type by Abuse Category ({p.capitalize()}, {stage_label})",
            f"fig{fig_n}_{prefix}_heatmap_{p}"))
        fig_n += 1

    if HAVE_WORDCLOUD:
        for p in PERSONA_ORDER:
            sub = [r for r in data if r["persona"] == p]
            text = " ".join(clean_text(r["raw_text"]).lower() for r in sub if r["raw_text"])
            path = wordcloud_image(text, f"{p.capitalize()} Persona ({stage_label})",
                                     f"fig{fig_n}_{prefix}_wordcloud_{p}")
            if path:
                embed(path)
            fig_n += 1

    return fig_n


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------

def infer_label(path, idx, override_labels):
    if override_labels:
        return override_labels[idx]
    base = os.path.basename(path).lower()
    m = re.search(r"condition([abc])", base)
    if m:
        return m.group(1).upper()
    return chr(ord("A") + idx)


def main():
    raw_args = sys.argv[1:]
    override_labels = None
    if "--labels" in raw_args:
        i = raw_args.index("--labels")
        main_args = raw_args[:i]
        override_labels = raw_args[i + 1:]
    else:
        main_args = raw_args

    if len(main_args) < 4:
        print("Usage: python3 build_stage_results.py <stageA.csv> <stageB.csv> <stageC.csv> <output.xlsx> "
              "[--labels \"Label A\" \"Label B\" \"Label C\"]")
        sys.exit(1)
    paths, output_file = main_args[:-1], main_args[-1]

    wb = Workbook()
    wb.remove(wb.active)

    ws_notes = wb.create_sheet("Notes")
    ws_notes.column_dimensions["A"].width = 110
    notes = [
        "Per-stage results for the Policy A/B/C ablation (run_llm_judge.py output), in the same seaborn-darkgrid "
        "chart style and green heat-mapped table style as the main persona results workbook (build_results.py).",
        "",
        "Rows with a judge JSON-parse failure (llm_parse_error non-empty) are dropped before any percentage is",
        "calculated. 'combined' = girlfriend + boyfriend rows pooled, same convention as the main workbook.",
    ]
    for i, line in enumerate(notes, start=1):
        c = ws_notes.cell(row=i, column=1, value=line)
        c.font = TITLE_FONT if i == 1 else NOTE_FONT

    fig_n = 1
    summary_rows = []
    for idx, path in enumerate(paths):
        label = infer_label(path, idx, override_labels)
        print(f"Loading stage {label}: {path}")
        data, dropped = load_stage(path)
        print(f"  {len(data) - sum(1 for d in data if d['persona']=='combined')} rows loaded, {dropped} dropped")
        fig_n = build_stage(wb, label, data, dropped, fig_n)
        for row in layer2_by_persona(data):
            summary_rows.append([label] + row)

    ws_sum = wb.create_sheet("Cross_stage_summary", 1)
    write_table_heatmap(ws_sum, ["Stage", "Persona", "n valid", "n excluded", "Harmful %", "Non-harmful %"],
                          summary_rows, title="Harmful % by persona, across all stages (for quick comparison)",
                          label_cols=2)
    autosize(ws_sum)

    order = ["Notes", "Cross_stage_summary"] + [s for s in wb.sheetnames if s not in ("Notes", "Cross_stage_summary")]
    wb._sheets = [wb[name] for name in order]

    wb.save(output_file)
    print(f"\nSaved {output_file}")
    print(f"Standalone PNG figures in {FIG_DIR}/")


if __name__ == "__main__":
    main()
